"""Regression tests for the v1.1.0 field-report bugs (CHANGELOG 1.1.1). No browser needed.

    python tests/test_regressions.py      (needs Pillow + python-pptx; exits non-zero on failure)

Each test builds on a temp copy of the demo workspace and its already-captured shots.
"""
import contextlib, io, json, os, pathlib, shutil, subprocess, sys, tempfile
from types import SimpleNamespace
from unittest import mock

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = pathlib.Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "plugin/skills/mslides/scripts"
sys.path.insert(0, str(SCRIPTS))
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches
import build_manual as B
import manual as M
import qa_render

TMP = pathlib.Path(tempfile.mkdtemp())


def workspace(name):
    """Same relative layout as the repo, so manual.json's relative template path resolves."""
    root = TMP / name
    shutil.copytree(ROOT / "examples", root / "examples")
    shutil.copytree(ROOT / "plugin/skills/mslides/templates", root / "plugin/skills/mslides/templates")
    return root / "examples/demo"


def build_task(ws, item):
    """Build one task slide in-process and return (pic, boxes, badges) from the saved deck."""
    cfg = json.loads((ws / "manual.json").read_text(encoding="utf-8"))
    out = ws / "out" / "one.pptx"
    B.init(cfg, ws, out)
    B.task_slide(item)
    with contextlib.redirect_stdout(io.StringIO()):
        B.finish()
    s = Presentation(out).slides[0]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    boxes = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    badges = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.auto_shape_type == MSO_SHAPE.OVAL
              and sh.left > pic.left]  # the numbered ovals of the step list sit left of the picture
    return pic, boxes, badges


def centre(sh):
    return sh.left + sh.width / 2, sh.top + sh.height / 2


def item_for(ws, shot, marks, steps=None):
    """A manifest-like item over a copy of an existing demo picture with its own marks."""
    shutil.copy(ws / "shots/viewer-01-list.png", ws / "shots" / f"{shot}.png")
    (ws / "shots" / f"{shot}.json").write_text(json.dumps({"w": 1600, "h": 900, "marks": marks}), encoding="utf-8")
    return {"chapter": "viewer", "task": "t", "kicker": "VIEWER · x", "shot": shot, "steps": steps or [f"step {i + 1}" for i in range(len(marks))]}


def test_diff_shots_badge_only():
    ws = workspace("diff")
    shutil.copytree(ws / "shots", ws / "shots-new")
    path = ws / "shots-new/viewer-01-list.json"
    meta = json.loads(path.read_text(encoding="utf-8"))
    meta["marks"][2]["badge"] = "left"  # picture and geometry identical: only the hint differs
    path.write_text(json.dumps(meta), encoding="utf-8")
    diff = lambda *a: subprocess.run([sys.executable, str(SCRIPTS / "diff_shots.py"), str(ws), *a], capture_output=True, text=True)
    r = diff()
    assert "changed   viewer-01-list" in r.stdout and "hints" in r.stdout, f"badge-only edit not reported:\n{r.stdout}{r.stderr}"
    assert "1 changed" in r.stdout and "5 unchanged" in r.stdout, f"other shots flagged:\n{r.stdout}"
    r = diff("--accept-changed")
    assert r.returncode == 0, r.stderr
    got = json.loads((ws / "shots/viewer-01-list.json").read_text(encoding="utf-8"))
    assert got["marks"][2].get("badge") == "left", "--accept-changed did not copy the badge edit into shots/"
    assert "0 changed" in diff().stdout, "after accepting, the shot must read unchanged"
    print("ok: diff_shots reports and accepts a badge-only edit")


def test_clamp_box_unit():
    inset = Inches(0.05)
    assert B._clamp_box(-100, -100, 5000000, 3000000, 0, 0, 4000000, 2000000, inset) == (inset, inset, 4000000 - 2 * inset, 2000000 - 2 * inset)
    assert B._clamp_box(500000, 500000, 100, 100, 0, 0, 4000000, 2000000, inset) == (500000, 500000, 100, 100), "inside box must not change"
    print("ok: _clamp_box")


def test_edge_marks_stay_inside_picture():
    ws = workspace("edge")
    # A drawer header at the right edge, a top bar at the top-left, a footer at the bottom edge.
    marks = [{"x": 1200, "y": 0, "w": 400, "h": 60}, {"x": 0, "y": 0, "w": 300, "h": 60}, {"x": 400, "y": 820, "w": 1200, "h": 80}]
    pic, boxes, _ = build_task(ws, item_for(ws, "edge-marks", marks))
    assert len(boxes) == 3
    for b in boxes:
        assert b.left >= pic.left and b.top >= pic.top and b.left + b.width <= pic.left + pic.width \
            and b.top + b.height <= pic.top + pic.height, "box outline spills outside the picture"
    print("ok: boxes touching the screenshot edge stay inside the picture")


def test_badge_hints_and_ambiguity_warning():
    ws = workspace("badge")
    box = {"x": 700, "y": 400, "w": 200, "h": 40}
    for hint in ("left", "right", "above", "below"):
        it = item_for(ws, f"hint-{hint}", [dict(box, badge=hint)])
        _, boxes, badges = build_task(ws, it)
        (bx, by), (cx, cy) = centre(boxes[0]), centre(badges[0])
        got = {"left": cx < boxes[0].left, "right": cx > boxes[0].left + boxes[0].width,
               "above": cy < boxes[0].top, "below": cy > boxes[0].top + boxes[0].height}[hint]
        assert got, f"badge hint {hint!r} put the badge at {cx - bx:.0f},{cy - by:.0f} from the box centre"
    try:
        build_task(ws, item_for(ws, "hint-bad", [dict(box, badge="diagonal")]))
    except SystemExit as e:
        assert "diagonal" in str(e) and "above" in str(e)
    else:
        raise AssertionError("an unknown badge hint must fail the build")
    # Two buttons 4 px apart; forcing mark 2's badge to the left puts it in the gap — the build must say so.
    adj = [{"x": 600, "y": 400, "w": 120, "h": 36}, {"x": 724, "y": 400, "w": 120, "h": 36, "badge": "left"}]
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        build_task(ws, item_for(ws, "hint-adjacent", adj))
    assert "badge 2 is about as close to box 1" in err.getvalue(), f"no ambiguity warning:\n{err.getvalue()}"
    free = [adj[0], {**adj[1], "badge": "right"}]
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        build_task(ws, item_for(ws, "hint-free", free))
    assert "about as close" not in err.getvalue(), f"false warning:\n{err.getvalue()}"
    print("ok: badge hints right/above/below, unknown hint fails, ambiguous spot warns")


def test_keynote_locked_screen():
    fake = lambda err: SimpleNamespace(returncode=1, stdout="", stderr=err)
    with mock.patch.object(M.R, "has_keynote", return_value=True), \
            mock.patch.object(M.subprocess, "run", return_value=fake("execution error: AppleEvent timed out. (-1712)")):
        try:
            M.render_pdf(TMP / "x.pptx")
        except SystemExit as e:
            assert "screen is locked" in str(e), e
        else:
            raise AssertionError("-1712 must stop the build")
    with mock.patch.object(M.R, "has_keynote", return_value=True), \
            mock.patch.object(M.subprocess, "run", return_value=fake("boom")):
        try:
            M.render_pdf(TMP / "x.pptx")
        except SystemExit as e:
            assert "boom" in str(e) and "locked" not in str(e)
        else:
            raise AssertionError("a Keynote failure must stop the build")
    print("ok: Keynote -1712 gives the locked-screen message")


def test_publish_keeps_pdfs_of_live_editions():
    ws = workspace("publish")
    cfg = json.loads((ws / "manual.json").read_text(encoding="utf-8"))
    live = M.edition_names(cfg, ["all"] + [c["code"] for c in cfg["chapters"]])
    out, staged = ws / "out", ws / "out/.staging"
    staged.mkdir(parents=True)
    for n in live + ["Demo-Manual-Gone"]:
        for ext in (".pptx", ".pdf", ".md"):
            (out / f"{n}{ext}").write_text("old", encoding="utf-8")
    (out / ".version").write_text(str(cfg["version"]), encoding="utf-8")
    for n in live:  # a build WITHOUT --pdf / --md: decks only
        (staged / f"{n}.pptx").write_text("new", encoding="utf-8")
    err = io.StringIO()
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(err):
        M.publish(cfg, ws, staged)
    for n in live:
        assert (out / f"{n}.pptx").read_text(encoding="utf-8") == "new"
        assert (out / f"{n}.pdf").exists(), f"{n}.pdf was archived by a build that merely skipped --pdf"
        assert (out / f"{n}.md").exists(), f"{n}.md was archived by a build that merely skipped --md"
    assert not list((out / "archive/removed").glob("Demo-Manual-ALL*")), "live edition archived as removed"
    assert (out / "archive/removed/Demo-Manual-Gone.pdf").exists(), "a really removed edition must still be archived"
    assert not (out / "Demo-Manual-Gone.pptx").exists()
    assert "Demo-Manual-ALL.pdf" in err.getvalue() and "may not match" in err.getvalue(), "stale PDFs must be flagged"
    print("ok: publish keeps PDFs/Markdown of live editions and flags them stale")


def test_qa_render_dpi():
    if not qa_render.pdftoppm():
        print("skip: qa_render (no pdftoppm)")
        return
    from PIL import Image
    pdf = TMP / "qa.pdf"
    Image.new("RGB", (600, 340), "white").save(pdf, resolution=72)  # 600x340 pt page
    pages = qa_render.render(pdf, TMP / "qa", dpi=200)
    assert len(pages) == 1
    w, h = Image.open(pages[0]).size
    assert abs(w - 600 * 200 / 72) <= 2 and abs(h - 340 * 200 / 72) <= 2, f"{w}x{h} is not 200 dpi"
    print(f"ok: qa_render at 200 dpi → {w}x{h}")


failed = []
try:
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
            except (AssertionError, SystemExit, Exception) as e:  # report every failing test, not just the first
                failed.append(name)
                print(f"FAIL {name}: {type(e).__name__}: {str(e)[:300]}")
finally:
    shutil.rmtree(TMP, ignore_errors=True)
if failed:
    sys.exit(f"{len(failed)} regression test(s) failed: {', '.join(failed)}")
print("REGRESSIONS OK")
