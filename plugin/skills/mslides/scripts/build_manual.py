"""Slide primitives for a user manual (.pptx) built on ANY template.

Layout grammar (one per task slide): kicker+title · numbered steps left · screenshot right with native callouts
(outline box + numbered oval per step) · optional tip in a fixed slot. Reference material is native tables.
Everything is a real PowerPoint shape so anyone can edit it in PowerPoint/Keynote; callouts are positioned from the element rects that
cap.mjs recorded at capture time (CSS px in the capture viewport), scaled onto the placed picture.

All template knowledge lives in manual.json → "template" (see references/template.md). Nothing here names a
template file, a layout, or a colour: a dark branded deck and a plain white one differ only in config.
Call init(cfg, workspace) once, then the slide functions, then finish().
"""
import copy, json, os, pathlib, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# Defaults suit a LIGHT template (the bundled templates/plain.pptx). A dark template must override at least
# "text"/"muted"/"tip_fill" (dark-on-dark text is the classic failure — it validates fine, only QA by eye sees it).
DEFAULT_COLORS = {
    "accent": "E5484D", "on_accent": "FFFFFF", "text": "1F2328", "muted": "4A4F57",
    "tip_fill": "FDECEC", "table_head": "E5484D", "table_row": "F4F5F7", "table_ink": "1F2328", "pic_line": "C8CCD2",
}
# The dark counterpart (templates/dark.pptx, background 1B1F24): inspect_template.py --suggest emits these for a
# dark theme, and references/template.md shows the same block — keep the three in step.
DARK_COLORS = {
    "text": "E8EAED", "muted": "B0B6BE", "tip_fill": "2A3038", "table_row": "262B32", "table_ink": "E8EAED", "pic_line": "5A6068",
}
# Content area [x, y, w, h] in inches: everything below the title and above the footer/logo.
# Measured on a 13.33" x 7.5" template: steps start at 2.0", tip slot ends at 6.5", footer/logo at ~6.8".
DEFAULT_AREA = [0.35, 1.95, 12.65, 4.8]

prs = SHOTS = OUT = C = None
AREA, FONT, SCRIPT_FONT, TIP_LABEL = DEFAULT_AREA, "Arial", None, "Tip: "
LAYOUT, COVER, OLD_IDS, FOOTER_SPS, TITLE_BOX = {}, {}, [], [], None


def _rgb(hex6):
    return RGBColor.from_string(hex6.lstrip("#").upper())


def _layout(ref):
    """{"layout": name, "master": i} or {"slide": n} (1-based: use that template slide's layout).
    By-slide exists because some templates repeat layout names across masters (e.g. "Title Slide" in three masters)."""
    if "slide" in ref:
        return prs.slides[ref["slide"] - 1].slide_layout
    master = prs.slide_masters[ref.get("master", 0)]
    for l in master.slide_layouts:
        if l.name == ref["layout"]:
            return l
    raise SystemExit(f"layout {ref['layout']!r} not in master {ref.get('master', 0)} — run inspect_template.py")


def init(cfg, workspace, out):
    global prs, SHOTS, OUT, C, AREA, FONT, SCRIPT_FONT, TIP_LABEL, LAYOUT, COVER, OLD_IDS, FOOTER_SPS, TITLE_BOX
    t = cfg["template"]
    tpl = pathlib.Path(os.path.expanduser(t["path"]))
    # Relative = relative to the workspace (<ws>/manual.json → "../templates/x.pptx"), so the same sources build
    # the same deck from any cwd — on a laptop or in CI.
    prs = Presentation(tpl if tpl.is_absolute() else workspace / tpl)
    SHOTS = workspace / cfg.get("shots", "shots")
    OUT = out
    C = {k: _rgb(v) for k, v in {**DEFAULT_COLORS, **t.get("colors", {})}.items()}
    AREA = t.get("area", DEFAULT_AREA)
    # script_font = complex-script face (<a:cs>) for languages the Latin font lacks, e.g. Thai → "Tahoma"
    # (Arial has no Thai glyphs; Tahoma ships with Office on Windows and macOS).
    FONT, SCRIPT_FONT = t.get("font", "Arial"), t.get("script_font")
    TIP_LABEL = t.get("tip_label", cfg.get("tip_label", "Tip: "))
    TITLE_BOX = t.get("title_box")
    LAYOUT = {k: _layout(t[k]) for k in ("cover", "divider", "content")}
    COVER = t["cover"]
    OLD_IDS = list(prs.slides._sldIdLst)  # template sample slides are dropped in finish()
    # Optional: some templates put the page footer on sample slides, not on the layout — copy it from one.
    f = t.get("footer")
    FOOTER_SPS = [copy.deepcopy(sh._element) for sh in prs.slides[f["slide"] - 1].shapes
                  if sh.is_placeholder and sh.placeholder_format.idx in f["placeholders"]] if f else []


# ---------- text helpers ----------
def run(p, text, size=None, bold=None, color=None):
    r = p.add_run()
    r.text = text
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if color is not None: r.font.color.rgb = color
    r.font.name = FONT
    if SCRIPT_FONT:
        rpr = r._r.get_or_add_rPr()
        cs = rpr.find(qn("a:cs"))
        if cs is None:  # not `or`: an lxml element with no children is falsy, which duplicated <a:cs>
            cs = etree.SubElement(rpr, qn("a:cs"))
        cs.set("typeface", SCRIPT_FONT)
    return r


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def rich(p, text, size, base=None, strong=None):
    """'**bold**' segments render in the strong colour + bold, the rest in the base colour."""
    base, strong = base or C["muted"], strong or C["text"]
    for i, part in enumerate(text.split("**")):
        if part:
            run(p, part, size=size, bold=bool(i % 2), color=strong if i % 2 else base)


def _ph(s, idx):
    return next((p for p in s.placeholders if p.placeholder_format.idx == idx), None)


# ---------- slide scaffolding ----------
def new_slide(kind, notes=""):
    s = prs.slides.add_slide(LAYOUT[kind])
    for el in FOOTER_SPS:
        sp = copy.deepcopy(el)
        # A copied footer keeps its template shape id, which can collide with a layout placeholder id → Keynote crashed.
        sp.find(".//" + qn("p:cNvPr")).set("id", str(s.shapes._next_shape_id))
        s.shapes._spTree.append(sp)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def heading(s, kicker, title):
    ph = _ph(s, 0)
    if ph is not None and not TITLE_BOX:
        p = ph.text_frame.paragraphs[0]
    else:  # layout has no title placeholder (or config overrides it): draw a text box where the template expects one
        p = tbox(s, *(TITLE_BOX or [AREA[0], 0.3, AREA[2], 1.2]), anchor=MSO_ANCHOR.BOTTOM).paragraphs[0]
    k = run(p, kicker.rstrip(": "), size=12, bold=True, color=C["accent"])  # a trailing ":" read as noise in QA
    k._r.get_or_add_rPr().set("spc", "150")
    etree.SubElement(p._p, qn("a:br"))
    run(p, title, size=30, bold=True, color=C["text"])


def title_slide(title, subtitle, notes=""):
    s = new_slide("cover", notes)
    t, sub = _ph(s, 0), _ph(s, 1)
    if t is None or sub is None:  # cover layout without placeholders: boxes from config
        tb = COVER.get("title_box", [1.0, 2.6, prs.slide_width / 914400 - 2.0, 1.2])
        sb = COVER.get("subtitle_box", [tb[0], tb[1] + tb[3] + 0.1, tb[2], 0.8])
        run(tbox(s, *tb).paragraphs[0], title, size=40, bold=True, color=C["text"])
        run(tbox(s, *sb).paragraphs[0], subtitle, size=20, color=C["muted"])
        return s
    if "title_left" in COVER:  # move both placeholders' left edge (e.g. to clear logo art on the left)
        for ph in (t, sub):
            left, top, width, height = ph.left, ph.top, ph.width, ph.height
            ph.left, ph.top, ph.height = Inches(COVER["title_left"]), top, height  # set all four: a partial xfrm zeroes the rest
            ph.width = left + width - Inches(COVER["title_left"])
    run(t.text_frame.paragraphs[0], title)
    run(sub.text_frame.paragraphs[0], subtitle)
    return s


def list_slide(title, items, notes=""):
    """Agenda / chapter-divider: title + list of (text, accent number shown before it)."""
    s = new_slide("divider", notes)
    t, body = _ph(s, 0), _ph(s, 1)
    # No placeholder: reuse the content slides' title box, which is already known to clear the template's logo.
    tp = t.text_frame.paragraphs[0] if t is not None else \
        tbox(s, *(TITLE_BOX or [AREA[0], 0.3, AREA[2], 1.2]), anchor=MSO_ANCHOR.BOTTOM).paragraphs[0]
    run(tp, title, **({} if t is not None else {"size": 36, "bold": True, "color": C["text"]}))
    tf = body.text_frame if body is not None else tbox(s, AREA[0], AREA[1], AREA[2], AREA[3])
    for i, (txt, n) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        # Number first ("1  Viewer"): a trailing accent number read as a stray page count in an outside QA run.
        if n: run(p, n + "  ", size=18, bold=True, color=C["accent"])
        run(p, txt, size=18, bold=True, **({} if body is not None else {"color": C["text"]}))
    return s


def number_badge(s, n, cx, cy, d=0.34):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx - Inches(d) / 2)), Emu(int(cy - Inches(d) / 2)), Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = C["accent"]
    o.line.color.rgb = C["on_accent"]; o.line.width = Pt(1.5)
    tf = o.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run(p, str(n), size=12, bold=True, color=C["on_accent"])
    return o


def bullets_slide(kicker, title, lines, notes=""):
    s = new_slide("content", notes)
    heading(s, kicker, title)
    w, h = AREA[2] - 0.65, AREA[3] - 0.8
    # 20pt lines at ~0.55 em per char, 1.2 line height, plus 14pt after each bullet: past the box the text
    # overflows onto the footer while the build "succeeds" — fail with a split instruction instead.
    # ponytail: char-count estimate tuned for Thai, conservative for Latin; measure real text metrics if it bites.
    per_line = max(1, int(w * 72 / (20 * 0.55)))
    need = sum(-(-len(f"{i + 1}.  {l}".replace("**", "")) // per_line) * 20 * 1.2 / 72 + 14 / 72 for i, l in enumerate(lines))
    if need > h:
        raise SystemExit(f"bullets {title!r} need ~{need:.1f}\" but the box is {h:.1f}\" — split into two entries or shorten")
    tf = tbox(s, AREA[0] + 0.25, AREA[1] + 0.25, w, h)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        rich(p, f"{i + 1}.  {line}", 20)
    return s


def crop_window(meta, min_frac=0.45, pad=60, img=None):
    """Smallest 16:9 window (≥ min_frac of the screen width) holding every mark plus padding, clamped to the image.
    Always 16:9 whatever the capture viewport: task_slide puts it in a 16:9 box and maps callouts with ONE scale
    factor, so a 4:3 window would stretch the picture and misplace every box vertically."""
    W, H = meta["w"], meta["h"]
    R = 16 / 9
    if not meta["marks"]:
        return 0, 0, W, H
    x0 = min(r["x"] for r in meta["marks"]) - pad
    y0 = min(r["y"] for r in meta["marks"]) - pad
    x1 = max(r["x"] + r["w"] for r in meta["marks"]) + pad
    y1 = max(r["y"] + r["h"] for r in meta["marks"]) + pad
    cw = max(x1 - x0, (y1 - y0) * R, W * min_frac)
    cw = min(cw, W, H * R)
    ch = cw / R
    cx = min(max((x0 + x1) / 2 - cw / 2, 0), W - cw)
    cy = min(max((y0 + y1) / 2 - ch / 2, 0), H - ch)
    # A very wide viewport caps the 16:9 window at H*R, which can cut off a mark at the far side; a callout drawn
    # outside the picture would land on the step text. Fail so the shot gets retaken with closer marks.
    def holds(x, y, w, h):
        return all(r["x"] >= x - 1 and r["x"] + r["w"] <= x + w + 1 and r["y"] >= y - 1 and r["y"] + r["h"] <= y + h + 1
                   for r in meta["marks"])

    if not holds(cx, cy, cw, ch):
        raise SystemExit(f"{meta.get('name', 'shot')}: marks span more than one 16:9 window — retake with marks closer together")
    if img is None:
        return cx, cy, cw, ch
    # The window is placed from mark geometry alone, so its edges can slice through UI text ("Demo Inventory" cropped
    # to "ntory", a heading cut in half — first outside QA run). Try slightly larger / shifted windows and keep the
    # one whose edges cross the least ink; ties go to the tightest (most zoomed) window, so readability wins.
    g = img.convert("L")
    if g.size != (W, H):
        g = g.resize((int(W), int(H)))
    px = g.load()

    def edge_ink(x, y, w, h):
        def line(pts):
            v = [px[min(int(a), W - 1), min(int(b), H - 1)] for a, b in pts]
            if not v:
                return 0.0
            bg = sorted(v)[len(v) // 2]
            return sum(abs(p - bg) > 40 for p in v) / len(v)
        n = 0.0
        if x > 1: n += line((x, y + i) for i in range(0, int(h), 2))
        if x + w < W - 1: n += line((x + w, y + i) for i in range(0, int(h), 2))
        if y > 1: n += line((x + i, y) for i in range(0, int(w), 2))
        if y + h < H - 1: n += line((x + i, y + h) for i in range(0, int(w), 2))
        return n

    # Zooming out is a cost, not a free escape: an edge on the image border reads as "no ink", so an unpenalised
    # search on a dense full-width screen just picks the whole screen and the UI text shrinks (most slides of a
    # real 93-slide manual did). Each 1% of extra width costs as much as a 1.5% edge crossing, and growth is capped at 12%.
    # Ties go to the least movement — ordering ties by coordinate nudged 40 calm crops up-left for nothing.
    base = edge_ink(cx, cy, cw, ch)
    best = (round(base, 2), 0.0, cx, cy, cw, ch)
    for grow in (1.0, 1.04, 1.08, 1.12):
        w2 = min(cw * grow, W, H * R)
        h2 = w2 / R
        for fx in (-0.06, -0.03, 0, 0.03, 0.06):
            for fy in (-0.06, -0.03, 0, 0.03, 0.06):
                x2 = min(max(cx + (cw - w2) / 2 + fx * w2, 0), W - w2)
                y2 = min(max(cy + (ch - h2) / 2 + fy * h2, 0), H - h2)
                # A moved window keeps 12 px around every mark (QA saw boxes 3 px from the edge) and never moves
                # its top edge down: page titles live at the top, and shifting down dropped them (3 slides of a real manual).
                if holds(x2 + 13, y2 + 13, w2 - 26, h2 - 26) and y2 <= cy + 0.5:
                    moved = abs(x2 - cx) / cw + abs(y2 - cy) / ch + (w2 / cw - 1)
                    best = min(best, (round(edge_ink(x2, y2, w2, h2) + 1.5 * (w2 / cw - 1), 2), moved, x2, y2, w2, h2))
    # Move only for a real gain (≥4% of one edge in ink): ink cannot tell a heading from a table row, so noise-level
    # gains just trade one partly-cut thing for another. QA on a real 93-slide manual: 13 better, 9 same, 1 borderline.
    return best[2:] if base - best[0] >= 0.04 else (cx, cy, cw, ch)


BADGE_HINTS = ("left", "right", "above", "below")
# A callout box is drawn `pad` outside its mark with a 2.25pt stroke centred on the edge: a mark that touches the
# screenshot frame would put the stroke on the slide background. Keep the whole box this far inside the picture.
FRAME_INSET = 0.05  # inches: pad 0.03 + half the stroke (0.016) + a hair


def _clamp_box(bx, by, bw, bh, left, top, width, height, inset):
    """Shrink box (bx, by, bw, bh) so it lies inside the rectangle (left, top, width, height) inset on all sides.
    Any unit, as long as all arguments share it. A box already inside is returned unchanged."""
    x0, y0 = max(bx, left + inset), max(by, top + inset)
    x1, y1 = min(bx + bw, left + width - inset), min(by + bh, top + height - inset)
    return x0, y0, max(x1 - x0, 0), max(y1 - y0, 0)


def task_slide(item):
    ax, ay, aw, ah = AREA
    s = new_slide("content", item.get("tip", ""))
    heading(s, item["kicker"], item["task"])
    # Steps (left column). The tip box has a FIXED slot at the bottom of the area so it can never slide down over
    # a template logo; when the steps would reach that slot, they shrink 14→12→11pt instead.
    TIP_TOP = ay + ah - 1.3
    # 0.25" guard: the char-count line estimate runs short for Thai at 12pt (a real deck's steps touched the tip).
    limit = TIP_TOP - 0.25 if item.get("tip") else ay + ah - 0.15

    def layout(size):
        per_line, line_h = {14: (38, 0.28), 12: (40, 0.25), 11: (44, 0.23)}[size]  # chars/line in 3.45" (Thai-tuned), in/line
        ys, y = [], ay + 0.05
        for step in item["steps"]:
            ys.append(y)
            y += line_h * max(1, round(len(step) / per_line + 0.49)) + 0.3
        return ys, y

    size = 14
    ys, end = layout(14)
    for smaller in (12, 11):
        if end <= limit:
            break
        size = smaller
        ys, end = layout(smaller)
    # `limit` keeps a 0.25" comfort guard and `end` includes the gap after the last step; the hard failure is the
    # last step's text actually reaching the tip slot (or the area bottom). A slide between the two (text ending
    # ~0.4" above the tip) passed several QA rounds — tight is fine, overlapping is not.
    hard = (TIP_TOP - 0.05) if item.get("tip") else ay + ah
    if end - 0.3 > hard:  # even 11pt runs into the tip slot / footer
        raise SystemExit(f"{item['shot']}: steps do not fit even at 11pt — shorten them or split the task into two slides")
    for i, (step, y) in enumerate(zip(item["steps"], ys), 1):
        number_badge(s, i, Inches(ax + 0.17), Inches(y + 0.17))
        tf = tbox(s, ax + 0.5, y, 3.45, 0.9)
        rich(tf.paragraphs[0], step, size)
    if item.get("tip"):
        tip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ax), Inches(TIP_TOP), Inches(3.95), Inches(1.05))
        tip.fill.solid(); tip.fill.fore_color.rgb = C["tip_fill"]
        tip.line.color.rgb = C["accent"]; tip.line.width = Pt(1)
        tf = tip.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.12); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT  # autoshape text defaults to centred
        tip_size = 12 if len(item["tip"]) <= 120 else 11  # 3 lines of 12pt fit the fixed slot; longer tips drop a point
        # Capacity of the 3.95" x 1.05" slot at 11pt (~0.55 em per char, 1.2 line height): beyond it the text
        # clips — fail and say so rather than ship an unreadable tip.
        per_line, max_lines = int((3.95 - 0.24) * 72 / (11 * 0.55)), int(1.05 * 72 / (11 * 1.2))
        if -(-len((TIP_LABEL + item["tip"]).replace("**", "")) // per_line) > max_lines:
            raise SystemExit(f"{item['shot']}: tip too long for its slot (~{per_line * max_lines} chars) — shorten it or move detail to notes")
        run(tf.paragraphs[0], TIP_LABEL, size=tip_size, bold=True, color=C["accent"])
        rich(tf.paragraphs[0], item["tip"], tip_size)
    # Screenshot (right column) + callouts. 16:9 box as wide as the area allows, shrunk if the area is too short.
    png = SHOTS / f"{item['shot']}.png"
    meta = {**json.loads((SHOTS / f"{item['shot']}.json").read_text(encoding="utf-8")), "name": item["shot"]}
    with Image.open(png) as shot_png:
        cx0, cy0, cw, ch = crop_window(meta, img=shot_png)
    box_w = min(aw - 4.2, ah * 16 / 9)
    BOX_W, BOX_H = Inches(box_w), Inches(box_w * 9 / 16)
    pic = s.shapes.add_picture(str(png), Inches(ax + aw - box_w), Inches(ay), width=BOX_W, height=BOX_H)
    # A full 1600px screen shrunk to ~8" makes UI text unreadable, so zoom into the marked area with a native crop
    # (an editor can "Reset crop" in PowerPoint to see the whole screen). The window keeps 16:9, so no distortion.
    pic.crop_left, pic.crop_top = cx0 / meta["w"], cy0 / meta["h"]
    pic.crop_right, pic.crop_bottom = 1 - (cx0 + cw) / meta["w"], 1 - (cy0 + ch) / meta["h"]
    pic.line.color.rgb = C["pic_line"]; pic.line.width = Pt(1)
    k = pic.width / cw
    D = Inches(0.3)
    pad = Inches(0.03)
    # 1) All boxes first, so no later box outline is drawn across an earlier badge (QA: digits cut).
    boxes = []
    for r in meta["marks"]:
        x, yy = pic.left + (r["x"] - cx0) * k, pic.top + (r["y"] - cy0) * k
        w, h = max(r["w"] * k, Inches(0.2)), max(r["h"] * k, Inches(0.16))
        bx, by, bw, bh = _clamp_box(x - pad, yy - pad, w + 2 * pad, h + 2 * pad,
                                    pic.left, pic.top, pic.width, pic.height, Inches(FRAME_INSET))
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(bx)), Emu(int(by)), Emu(int(bw)), Emu(int(bh)))
        box.fill.background(); box.line.color.rgb = C["accent"]; box.line.width = Pt(2.25)
        boxes.append((bx, by, bw, bh))

    # 2) Then each badge ENTIRELY outside its own box (half-inside badges hid button labels and values in QA).
    #    Candidates are scored by cost = (geometric hits, ink under the badge, listed order); lowest wins.
    off = D * 0.36  # centre offset from the corner along each axis → badge edge just kisses the corner

    def dist(cx, cy, b):  # centre → nearest point of box b (0 when inside)
        bx, by, bw, bh = b
        dx = max(bx - cx, 0, cx - (bx + bw)); dy = max(by - cy, 0, cy - (by + bh))
        return (dx * dx + dy * dy) ** 0.5

    AMBIGUITY = Inches(0.05)

    def rival(cx, cy, own):
        """Index (0-based) of a box the badge at (cx, cy) is about as close to as to its own, else None.
        Only a cost in hits(): when every candidate scores it, min() still picks one — so it is also reported."""
        mine = dist(cx, cy, boxes[own])
        near = [(dist(cx, cy, b), i) for i, b in enumerate(boxes) if i != own]
        d, i = min(near, default=(None, None))
        return i if d is not None and d < mine + AMBIGUITY else None

    def hits(cx, cy, own):
        r = D / 2
        n_hit = 0
        # Skip the badge's OWN box: a corner badge kisses its corner by design. Counting it made every corner
        # score 1 while an above-centred spot scored 0, so above-centre won even over text (seen in QA).
        mine = dist(cx, cy, boxes[own])
        for i, b in enumerate(boxes):
            if i == own:
                continue
            d = dist(cx, cy, b)
            n_hit += d < r  # overlaps another box
            # About as close to a neighbour as to its own box → reads as the neighbour's number (seen in QA).
            n_hit += d < mine + AMBIGUITY
        n_hit += 10 * sum(1 for px, py in placed if abs(cx - px) < D and abs(cy - py) < D)
        r = D / 2
        inside = (pic.left + r <= cx <= pic.left + pic.width - r) and (pic.top + r <= cy <= pic.top + pic.height - r)
        return n_hit + (0 if inside else 5)

    # Boxes say where the TARGET is, not where the TEXT is: a field's label sits just above it, so a corner that
    # is geometrically free can still land on a label (seen in QA). Measure it instead:
    # ink = share of pixels under the candidate badge that differ strongly from the local background.
    shot_img = Image.open(png).convert("L")

    def ink(cx, cy):
        sx = cx0 + (cx - pic.left) / k  # slide EMU → screenshot px
        sy = cy0 + (cy - pic.top) / k
        rpx = (D / 2) * 1.3 / k  # sample a little wider than the badge: a digit at its rim still gets covered
        patch = shot_img.crop((int(sx - rpx), int(sy - rpx), int(sx + rpx), int(sy + rpx)))
        pw, ph = patch.size
        px = list(patch.getdata()) or [255]
        bg = sorted(px)[len(px) // 2]  # median ≈ background grey level
        on = [[abs(px[y * pw + x] - bg) > 60 for x in range(pw)] for y in range(ph)]
        # A straight rule (panel edge, table border) is not text: covering it hides nothing, yet it scored like
        # text and pushed badges off a row's side onto an ambiguous corner (seen in QA). Drop full lines.
        rows = {y for y in range(ph) if sum(on[y]) > 0.8 * pw}
        cols = {x for x in range(pw) if sum(on[y][x] for y in range(ph)) > 0.8 * ph}
        n = sum(1 for y in range(ph) for x in range(pw) if on[y][x] and y not in rows and x not in cols)
        return n / max(1, pw * ph)

    placed = []
    for n, (bx, by, bw, bh) in enumerate(boxes, 1):
        gap = Inches(0.03)
        corners = [(bx + bw + off, by - off), (bx + bw + off, by + bh + off), (bx - off, by - off), (bx - off, by + bh + off)]
        # Side-middle: for boxes stacked with no gap (table rows) every corner sits on the
        # shared edge between two boxes; beside the row's middle it can only belong to that row.
        sides = [(bx - D / 2 - gap, by + bh / 2), (bx + bw + D / 2 + gap, by + bh / 2)]
        if bh < Inches(0.45):
            # Buttons/chips sit in ROWS: a corner badge lands in the gap between two neighbours and reads as the
            # neighbour's number (seen in QA). Centred directly above (or below) its own button
            # it can only belong to that button — try those first.
            corners = [(bx + bw / 2, by - D / 2 - gap), (bx + bw / 2, by + bh + D / 2 + gap)] + sides + corners
        else:
            corners = sides + corners  # beside the middle is unambiguous even for rows stacked with no gap

        # Geometric collisions dominate; then least ink, even 1% — a digit at the badge's rim reads as ~1%
        # (above-centre once covered the first digit of an amount); ties keep the listed order.
        def cost(c):
            i = ink(*c)
            return (hits(*c, n - 1), round(i, 2), corners.index(c))
        hint = meta["marks"][n - 1].get("badge")  # per-mark override where ink can't tell text from chrome
        if hint:
            if hint not in BADGE_HINTS:
                raise SystemExit(f"{item['shot']}: mark {n}: unknown badge hint {hint!r} — use one of {', '.join(BADGE_HINTS)}")
            above, below = (bx + bw / 2, by - D / 2 - gap), (bx + bw / 2, by + bh + D / 2 + gap)
            corners = [{"left": sides[0], "right": sides[1], "above": above, "below": below}[hint]]
        cx, cy = min(corners, key=cost)
        other = rival(cx, cy, n - 1)
        if other is not None:
            print(f"warning: {item['shot']}: badge {n} is about as close to box {other + 1} as to its own — "
                  f"set \"badge\" on mark {n} to {'/'.join(BADGE_HINTS)}, or separate the boxes", file=sys.stderr)
        if os.environ.get("BADGE_DEBUG") == item["shot"]:
            print(n, [(corners.index(c), cost(c), round(ink(*c), 3)) for c in corners])
        cx = max(pic.left + D / 2, min(cx, pic.left + pic.width - D / 2))
        cy = max(pic.top + D / 2, min(cy, pic.top + pic.height - D / 2))
        placed.append((cx, cy))
        number_badge(s, n, cx, cy, d=0.3)
    return s


def table_slide(kicker, title, header, rows, widths, notes="", size=11):
    s = new_slide("content", notes)
    heading(s, kicker, title)
    total_w = sum(widths)
    if total_w > AREA[2] + 0.01:  # wider than the area runs off the slide edge — no silent overflow
        raise SystemExit(f"table {title!r}: widths sum to {total_w}\" but the content area is {AREA[2]}\" wide — narrow the columns")
    # Estimated height: per row, the most-wrapped cell (chars per line ≈ usable width / ~0.55 em per char),
    # at ~1.25 line height plus cell margins. A table that would run past the area overlaps the footer — fail and
    # say to split it, rather than ship a clipped reference slide.
    def lines(text, w):
        per_line = max(1, int((w - 0.16) * 72 / (size * 0.55)))
        return max(1, -(-len(text.replace("**", "")) // per_line))
    est = sum(max(lines(str(v), widths[j]) for j, v in enumerate(r)) * size * 1.25 / 72 + 0.08 for r in [header] + rows)
    if est > AREA[3] + 0.2:
        raise SystemExit(f"table {title!r} needs ~{est:.1f}\" but the content area is {AREA[3]}\" — split it into two "
                         f"entries in manual.json (e.g. '… (1/2)', '… (2/2)') or shorten the cells")
    gt = s.shapes.add_table(len(rows) + 1, len(header), Inches(AREA[0]), Inches(AREA[1]), Inches(total_w),
                            Inches(0.3 * (len(rows) + 1)))
    t = gt.table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, row in enumerate([header] + rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = C["table_head"] if i == 0 else C["table_row"]
            c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            if i == 0:
                run(p, val, size=size, bold=True, color=C["on_accent"])
            else:
                rich(p, val, size, base=C["table_ink"], strong=C["table_ink"])
                if j == 0 and p.runs: p.runs[0].font.bold = True  # a blank grouping cell has no run
    return s


def finish():
    for sid in OLD_IDS:
        prs.part.drop_rel(sid.get(qn("r:id")))
        prs.slides._sldIdLst.remove(sid)
    ext_lst = prs.part._element.find(qn("p:extLst"))  # template sections name deleted slide ids -> repair prompt
    if ext_lst is not None:
        for ext in list(ext_lst):
            if any(ch.tag.endswith("}sectionLst") for ch in ext):
                ext_lst.remove(ext)
    # A template without a notes master gets one from python-pptx on the first speaker note, related to the
    # presentation but NOT listed in <p:notesMasterIdLst>. PowerPoint shrugs; Keynote refuses the whole file as
    # "invalid format". List it.
    root = prs.part._element
    nm = next((r for r in prs.part.rels.values() if r.reltype.endswith("/notesMaster")), None)
    if nm is not None and root.find(qn("p:notesMasterIdLst")) is None:
        lst = etree.Element(qn("p:notesMasterIdLst"))
        etree.SubElement(lst, qn("p:notesMasterId")).set(qn("r:id"), nm.rId)
        root.find(qn("p:sldMasterIdLst")).addnext(lst)  # schema order: sldMasterIdLst, notesMasterIdLst, …
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT, len(prs.slides), "slides")
    return OUT
