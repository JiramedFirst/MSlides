"""Cut a large template down to what one manual uses, so it can live in version control.

  python slim_template.py <template.pptx> <ws>/manual.json <out.pptx>

Keeps only the layouts manual.json's "template" block refers to (cover/divider/content, plus the layouts of any
slide it refers to) and only the referenced sample slides (e.g. the footer source), then rewrites manual.json:
"path" → <out.pptx> (relative to the workspace) and every {"slide": n} → its new number. python-pptx writes only
the parts still reachable, so the media of dropped layouts and sample slides disappear (a 17 MB template → ~0.5 MB).
"""
import json, os, pathlib, sys
from pptx import Presentation
from pptx.oxml.ns import qn

src, cfg_path, out = pathlib.Path(sys.argv[1]).expanduser(), pathlib.Path(sys.argv[2]).resolve(), pathlib.Path(sys.argv[3]).resolve()
cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
t = cfg["template"]
prs = Presentation(src)
slides = list(prs.slides)

# Which sample slides and layouts does the config reach?
refs = [t[k] for k in ("cover", "divider", "content")] + ([t["footer"]] if t.get("footer") else [])
keep_slides = sorted({r["slide"] for r in refs if "slide" in r})
keep_layouts = {id(slides[n - 1].slide_layout) for n in keep_slides}
for r in refs:
    if "layout" in r:
        master = prs.slide_masters[r.get("master", 0)]
        keep_layouts |= {id(l) for l in master.slide_layouts if l.name == r["layout"]}

# Drop every other sample slide (same mechanics as build_manual.finish)…
for i, sid in reversed(list(enumerate(prs.slides._sldIdLst))):
    if i + 1 not in keep_slides:
        prs.part.drop_rel(sid.get(qn("r:id")))
        prs.slides._sldIdLst.remove(sid)
# …then every layout nothing uses any more (a layout still used by a kept slide stays by construction).
for master in prs.slide_masters:
    for layout in list(master.slide_layouts):
        if id(layout) not in keep_layouts:
            master.slide_layouts.remove(layout)
ext = prs.part._element.find(qn("p:extLst"))  # section list names deleted slides → PowerPoint repair prompt
if ext is not None:
    for e in list(ext):
        if any(ch.tag.endswith("}sectionLst") for ch in e):
            ext.remove(e)
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(out)

# Remap {"slide": n} to the kept slides' new 1-based positions and point "path" at the slim copy.
new_no = {old: i + 1 for i, old in enumerate(keep_slides)}
for r in refs:
    if "slide" in r:
        r["slide"] = new_no[r["slide"]]
t["path"] = os.path.relpath(out, cfg_path.parent)
cfg_path.write_text(json.dumps(cfg, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print(f"{src.stat().st_size / 1e6:.1f} MB → {out.stat().st_size / 1e6:.2f} MB ({len(keep_slides)} slide(s) kept); "
      f"manual.json template.path = {t['path']}")
